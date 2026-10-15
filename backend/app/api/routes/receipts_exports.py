from __future__ import annotations

import io
from datetime import datetime, timezone
from decimal import Decimal
from xml.sax.saxutils import escape

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import select, text
from sqlalchemy.orm import Session

from app.api.deps import get_current_profile, get_profile_roles, require_roles
from app.db.session import get_db
from app.models import AuditLog, Bill, BillItem, Patient, Payment, Profile, Receipt

router = APIRouter(tags=["receipts and spreadsheet exports"])
STAFF_BILL_ROLES = {"administrator", "cashier", "information_officer"}


def _can_read_bill(db: Session, profile: Profile, bill: Bill) -> None:
    roles = set(get_profile_roles(profile.id, db))
    if roles.intersection(STAFF_BILL_ROLES):
        return
    if "patient" in roles:
        patient_id = db.scalar(select(Patient.id).where(Patient.patient_profile_id == profile.id))
        if patient_id == bill.patient_id:
            return
    raise HTTPException(status_code=404, detail="Receipt not found.")


@router.get("/receipts/{receipt_number}/pdf", summary="Download a payment receipt as PDF")
def download_receipt_pdf(
    receipt_number: str,
    profile: Profile = Depends(get_current_profile),
    db: Session = Depends(get_db),
) -> StreamingResponse:
    record = db.execute(
        select(Receipt, Payment, Bill, Patient)
        .join(Payment, Payment.id == Receipt.payment_id)
        .join(Bill, Bill.id == Payment.bill_id)
        .join(Patient, Patient.id == Bill.patient_id)
        .where(Receipt.receipt_number == receipt_number)
    ).first()
    if not record:
        raise HTTPException(status_code=404, detail="Receipt not found.")
    receipt, payment, bill, patient = record
    _can_read_bill(db, profile, bill)
    items = db.scalars(select(BillItem).where(BillItem.bill_id == bill.id).order_by(BillItem.id)).all()

    output = io.BytesIO()
    document = SimpleDocTemplate(output, pagesize=letter, rightMargin=0.65*inch,
                                 leftMargin=0.65*inch, topMargin=0.55*inch, bottomMargin=0.55*inch)
    styles = getSampleStyleSheet()
    story = [
        Paragraph("NexusCare Hospital", styles["Title"]),
        Paragraph("OFFICIAL PAYMENT RECEIPT", styles["Heading2"]),
        Spacer(1, 10),
        Paragraph(f"<b>Receipt No:</b> {escape(receipt.receipt_number)}", styles["Normal"]),
        Paragraph(f"<b>Date:</b> {receipt.issued_at:%Y-%m-%d %H:%M %Z}" if receipt.issued_at else "Date: —", styles["Normal"]),
        Paragraph(f"<b>Patient:</b> {escape(patient.first_name)} {escape(patient.last_name)}", styles["Normal"]),
        Paragraph(f"<b>Medical Record No:</b> {escape(patient.medical_record_number)}", styles["Normal"]),
        Paragraph(f"<b>Bill No:</b> {escape(bill.bill_number)}", styles["Normal"]),
        Spacer(1, 12),
    ]
    data = [["Description", "Qty", "Unit price (PHP)", "Amount (PHP)"]]
    for item in items:
        data.append([Paragraph(escape(item.description), styles["Normal"]), str(item.quantity), f"{item.unit_price:,.2f}", f"{item.line_total:,.2f}"])
    data.extend([
        ["", "", "Bill total", f"{bill.total:,.2f}"],
        ["", "", "This payment", f"{payment.amount:,.2f}"],
        ["", "", "Method / status", f"{payment.payment_method} / {payment.status}"],
    ])
    table = Table(data, colWidths=[3.1*inch, 0.5*inch, 1.25*inch, 1.25*inch], repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#16324f")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#d4dbe3")),
        ("ALIGN", (1, 1), (-1, -1), "RIGHT"),
        ("BACKGROUND", (0, -3), (-1, -1), colors.HexColor("#f2f5f8")),
        ("PADDING", (0, 0), (-1, -1), 7),
    ]))
    story.extend([table, Spacer(1, 18),
                  Paragraph("This receipt confirms the payment recorded in NexusCare. Keep it for your records.",
                            styles["Italic"])])
    document.build(story)
    db.add(AuditLog(actor_profile_id=profile.id, action="receipt.pdf_downloaded",
                    resource_type="receipt", resource_id=str(receipt.id),
                    details={"receipt_number": receipt.receipt_number}))
    db.commit()
    output.seek(0)
    return StreamingResponse(output, media_type="application/pdf", headers={
        "Content-Disposition": f'attachment; filename="receipt-{receipt.receipt_number}.pdf"'
    })


def _styled_sheet(workbook: Workbook, title: str, headers: list[str]):
    sheet = workbook.active
    sheet.title = title
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="16324F")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    return sheet


def _xlsx_response(workbook: Workbook, name: str) -> StreamingResponse:
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = f"nexuscare-{name}-{datetime.now(timezone.utc):%Y%m%d-%H%M%S}.xlsx"
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                            headers={"Content-Disposition": f'attachment; filename="{filename}"'})


@router.get("/exports/bills.xlsx", summary="Export bills to Excel")
def export_bills_xlsx(
    bill_status: str | None = Query(default=None, alias="status", pattern="^(open|partially_paid|paid|void)$"),
    profile: Profile = Depends(require_roles("administrator", "cashier", "information_officer")),
    db: Session = Depends(get_db),
) -> StreamingResponse:
    rows = db.execute(text("""
        select b.bill_number, b.patient_id, p.medical_record_number, p.first_name, p.last_name,
               b.status, b.subtotal, b.adjustments, b.total,
               coalesce(sum(case when pay.status='posted' then pay.amount else 0 end), 0) as paid,
               b.created_at
        from public.bills b join public.patients p on p.id=b.patient_id
        left join public.payments pay on pay.bill_id=b.id
        where (:bill_status is null or b.status=:bill_status)
        group by b.id, p.id order by b.created_at desc limit 5000
    """), {"bill_status": bill_status}).mappings().all()
    workbook = Workbook()
    sheet = _styled_sheet(workbook, "Bills", ["Bill number", "Patient ID", "MRN", "First name", "Last name",
        "Status", "Subtotal (PHP)", "Adjustments (PHP)", "Total (PHP)", "Paid (PHP)", "Balance (PHP)", "Created at"])
    for row in rows:
        total = Decimal(str(row["total"] or 0))
        paid = Decimal(str(row["paid"] or 0))
        sheet.append([row["bill_number"], str(row["patient_id"]), row["medical_record_number"],
            row["first_name"], row["last_name"], row["status"], float(row["subtotal"] or 0),
            float(row["adjustments"] or 0), float(total), float(paid), float(max(Decimal("0"), total-paid)),
            row["created_at"].isoformat() if row["created_at"] else ""])
    for col in "GHIJK":
        for cell in sheet[col][1:]:
            cell.number_format = '"₱"#,##0.00'
    db.add(AuditLog(actor_profile_id=profile.id, action="data.exported", resource_type="bills",
                    resource_id="xlsx", details={"row_count": len(rows), "status": bill_status}))
    db.execute(text("""insert into public.export_audit_events(actor_profile_id, export_type, format, filters, row_count)
        values (:actor, 'bills', 'xlsx', cast(:filters as jsonb), :count)"""),
        {"actor": profile.id, "filters": __import__("json").dumps({"status": bill_status}), "count": len(rows)})
    db.commit()
    return _xlsx_response(workbook, "bills")


@router.get("/exports/payments.xlsx", summary="Export payments to Excel")
def export_payments_xlsx(
    profile: Profile = Depends(require_roles("administrator", "cashier")),
    db: Session = Depends(get_db),
) -> StreamingResponse:
    rows = db.execute(text("""
        select pay.id, b.bill_number, b.patient_id, p.medical_record_number,
               pay.amount, pay.payment_method, pay.status, pay.receipt_number, pay.received_at
        from public.payments pay join public.bills b on b.id=pay.bill_id
        join public.patients p on p.id=b.patient_id
        order by pay.received_at desc limit 5000
    """)).mappings().all()
    workbook = Workbook()
    sheet = _styled_sheet(workbook, "Payments", ["Payment ID", "Bill number", "Patient ID", "MRN",
        "Amount (PHP)", "Method", "Status", "Receipt number", "Received at"])
    for row in rows:
        sheet.append([str(row["id"]), row["bill_number"], str(row["patient_id"]),
            row["medical_record_number"], float(row["amount"]), row["payment_method"], row["status"],
            row["receipt_number"], row["received_at"].isoformat() if row["received_at"] else ""])
    for cell in sheet["E"][1:]:
        cell.number_format = '"₱"#,##0.00'
    db.add(AuditLog(actor_profile_id=profile.id, action="data.exported", resource_type="payments",
                    resource_id="xlsx", details={"row_count": len(rows)}))
    db.execute(text("""insert into public.export_audit_events(actor_profile_id, export_type, format, filters, row_count)
        values (:actor, 'payments', 'xlsx', '{}'::jsonb, :count)"""),
        {"actor": profile.id, "count": len(rows)})
    db.commit()
    return _xlsx_response(workbook, "payments")
